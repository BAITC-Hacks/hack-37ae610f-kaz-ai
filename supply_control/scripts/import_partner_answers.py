from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from replenishment.partner_answers import (  # noqa: E402
    PartnerAnswerError,
    load_questionnaire,
    parse_answer_rows,
)


QUESTIONNAIRE = ROOT / "config/partner_questionnaire.json"
OUTPUT = ROOT / "config/partner_answers.json"


def read_answer_rows(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Вопросы" not in workbook.sheetnames:
            raise PartnerAnswerError("В книге отсутствует лист «Вопросы»")
        sheet = workbook["Вопросы"]
        header_row = None
        headers: dict[int, str] = {}
        for row in sheet.iter_rows(min_row=1, max_row=20):
            values = [str(cell.value or "").strip() for cell in row]
            if "ID" in values and "Ответ партнера" in values:
                header_row = row[0].row
                headers = {
                    index: value for index, value in enumerate(values) if value
                }
                break
        if header_row is None:
            raise PartnerAnswerError(
                "Не найдена строка заголовков с колонками ID и «Ответ партнера»"
            )
        required_headers = {
            "ID",
            "Ответ партнера",
            "Комментарий или источник партнера",
        }
        missing_headers = required_headers - set(headers.values())
        if missing_headers:
            raise PartnerAnswerError(
                "Отсутствуют колонки: " + ", ".join(sorted(missing_headers))
            )

        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(
            min_row=header_row + 1, max_row=500, values_only=True
        ):
            row = {header: values[index] for index, header in headers.items()}
            if not str(row.get("ID") or "").strip():
                continue
            rows.append(row)
        return rows
    finally:
        workbook.close()


def rebuild_outputs() -> None:
    commands = [
        [sys.executable, str(ROOT / "scripts/generate_demo_recommendations.py")],
        [sys.executable, str(ROOT / "scripts/build_integrated_site.py")],
    ]
    for command in commands:
        subprocess.run(command, cwd=ROOT, check=True)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Импорт ответов ТОО Электрокомплект из утвержденного XLSX-шаблона"
    )
    parser.add_argument("workbook", type=Path, help="Заполненный XLSX-файл партнера")
    parser.add_argument(
        "--output", type=Path, default=OUTPUT, help="Куда сохранить проверенный JSON"
    )
    parser.add_argument(
        "--rebuild",
        action="store_true",
        help="Пересчитать рекомендации и обновить локальный интерфейс",
    )
    args = parser.parse_args()

    questionnaire = load_questionnaire(QUESTIONNAIRE)
    answer_rows = read_answer_rows(args.workbook)
    payload = parse_answer_rows(answer_rows, questionnaire)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if args.rebuild:
        rebuild_outputs()
    print(
        json.dumps(
            {
                "output": str(args.output),
                "answered": payload["answered_count"],
                "questions": payload["question_count"],
                "complete": payload["complete"],
                "core_parameters_confirmed": payload["core_parameters_confirmed"],
                "missing_required_keys": payload["missing_required_keys"],
                "rebuilt": args.rebuild,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
