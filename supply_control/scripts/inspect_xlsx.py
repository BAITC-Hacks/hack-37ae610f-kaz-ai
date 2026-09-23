from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect_sheet(ws, preview_rows: int = 18, preview_cols: int = 35):
    previews = []
    nonempty_by_col = Counter()
    type_by_col: dict[int, Counter] = {}
    min_row = None
    max_row = 0
    min_col = None
    max_col = 0
    formula_count = 0

    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        row_preview = []
        row_has_value = False
        for col_idx, cell in enumerate(row, start=1):
            value = cell.value
            if value is not None:
                row_has_value = True
                nonempty_by_col[col_idx] += 1
                type_by_col.setdefault(col_idx, Counter())[type(value).__name__] += 1
                min_col = col_idx if min_col is None else min(min_col, col_idx)
                max_col = max(max_col, col_idx)
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
            if row_idx <= preview_rows and col_idx <= preview_cols:
                row_preview.append(clean(value))
        if row_has_value:
            min_row = row_idx if min_row is None else min(min_row, row_idx)
            max_row = row_idx
        if row_idx <= preview_rows:
            while row_preview and row_preview[-1] is None:
                row_preview.pop()
            previews.append(row_preview)

    return {
        "title": ws.title,
        "reported_max_row": ws.max_row,
        "reported_max_column": ws.max_column,
        "nonempty_bounds": {
            "min_row": min_row,
            "max_row": max_row,
            "min_col": min_col,
            "max_col": max_col,
        },
        "formula_count": formula_count,
        "nonempty_by_col": {str(k): v for k, v in nonempty_by_col.items()},
        "types_by_col": {
            str(k): dict(v) for k, v in type_by_col.items()
        },
        "preview": previews,
    }


def inspect_book(path: Path):
    wb = load_workbook(path, read_only=True, data_only=False)
    out = {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sheets": [],
    }
    for ws in wb.worksheets:
        out["sheets"].append(inspect_sheet(ws))
    wb.close()
    return out


def main():
    root = Path(sys.argv[1])
    output = Path(sys.argv[2])
    books = [inspect_book(path) for path in sorted(root.rglob("*.xlsx"))]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(books, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"files": len(books), "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
