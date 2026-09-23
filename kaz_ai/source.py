"""Read partner Excel exports without writing their contents into the repository."""

from __future__ import annotations

import io
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .engine import Arrival, Product, Sale


MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
    "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}


def _month(value: object) -> str | None:
    text = str(value or "").lower().strip()
    year = re.search(r"20\d{2}", text)
    if not year:
        return None
    for prefix, number in MONTHS.items():
        if text.startswith(prefix):
            return f"{year.group()}-{number:02d}"
    return None


def _number(value: object) -> float | None:
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    if isinstance(value, str):
        try:
            result = float(value.replace(" ", "").replace("\xa0", "").replace(",", "."))
            return result if math.isfinite(result) else None
        except ValueError:
            pass
    return None


def _code(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _decoded_name(info: zipfile.ZipInfo) -> str:
    try:
        return info.filename.encode("cp437").decode("cp866")
    except UnicodeEncodeError:
        return info.filename


def _sheet(archive: zipfile.ZipFile, info: zipfile.ZipInfo):
    workbook = load_workbook(io.BytesIO(archive.read(info)), read_only=True, data_only=True, keep_links=False)
    return workbook, workbook.worksheets[0]


def _find_columns(header: tuple, *names: str) -> int:
    for name in names:
        for index, value in enumerate(header):
            if str(value or "").strip().lower() == name.lower():
                return index
    raise ValueError(f"Missing column: {names}")


def _monthly(archive: zipfile.ZipFile, info: zipfile.ZipInfo, products: dict[str, Product], field: str) -> int:
    workbook, sheet = _sheet(archive, info)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    code_col = _find_columns(header, "Номенклатура.Код")
    name_col = _find_columns(header, "Номенклатура")
    month_cols = [(i, key) for i, value in enumerate(header) if (key := _month(value))]
    count = 0
    for row in rows:
        code = _code(row[code_col]) if code_col < len(row) else None
        if not code:
            continue
        product = products.setdefault(code, Product("", code, "", {}))
        if not product.name and name_col < len(row):
            product.name = str(row[name_col] or "").strip()
        target = product.monthly_sales if field == "sales" else product.opening_stock
        for col, month in month_cols:
            value = _number(row[col]) if col < len(row) else None
            if value is not None:
                target[month] = value
        count += 1
    workbook.close()
    return count


def _moq(archive: zipfile.ZipFile, info: zipfile.ZipInfo, products: dict[str, Product]) -> int:
    workbook, sheet = _sheet(archive, info)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    code_col = _find_columns(header, "Код 1с", "Номенклатура.Код")
    multiple_col = _find_columns(header, "Мин. разр. к отгр.", "Кратность")
    count = 0
    for row in rows:
        code = _code(row[code_col]) if code_col < len(row) else None
        if not code:
            continue
        multiple = _number(row[multiple_col]) if multiple_col < len(row) else None
        if multiple and multiple > 0:
            products.setdefault(code, Product("", code, "", {})).moq = max(1, math.ceil(multiple))
        count += 1
    workbook.close()
    return count


def _detail(archive: zipfile.ZipFile, info: zipfile.ZipInfo, products: dict[str, Product]) -> int:
    workbook, sheet = _sheet(archive, info)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    date_col = _find_columns(header, "Дата")
    code_col = _find_columns(header, "Код")
    qty_col = _find_columns(header, "Количество")
    doc_col = _find_columns(header, "Номер")
    customer_col = next((i for i, value in enumerate(header)
                         if "клиент" in str(value or "").lower()
                         and ("обезлич" in str(value).lower() or "аноним" in str(value).lower())), None)
    count = 0
    for row in rows:
        code = _code(row[code_col]) if code_col < len(row) else None
        qty = _number(row[qty_col]) if qty_col < len(row) else None
        if not code or qty is None or qty <= 0:
            continue
        value = row[date_col]
        try:
            day = value.date() if isinstance(value, datetime) else datetime.strptime(str(value), "%d.%m.%Y %H:%M:%S").date()
        except (ValueError, TypeError):
            continue
        document = _code(row[doc_col]) or f"row-{count}"
        customer = _code(row[customer_col]) if customer_col is not None and customer_col < len(row) else None
        products.setdefault(code, Product("", code, "", {})).sales.append(
            Sale(day.strftime("%Y-%m"), qty, document, customer)
        )
        count += 1
    workbook.close()
    return count


def _iek_transit(archive: zipfile.ZipFile, info: zipfile.ZipInfo, products: dict[str, Product]) -> int:
    workbook, sheet = _sheet(archive, info)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    code_col = _find_columns(header, "Код 1с")
    dated_cols: list[tuple[int, date]] = []
    for i, value in enumerate(header):
        match = re.search(r"поступление до (\d{2})\.(\d{2})\.(20\d{2})", str(value or ""))
        if match:
            dated_cols.append((i, date(int(match[3]), int(match[2]), int(match[1]))))
    count = 0
    for row in rows:
        code = _code(row[code_col]) if code_col < len(row) else None
        if not code:
            continue
        product = products.setdefault(code, Product("", code, "", {}))
        for i, due in dated_cols:
            qty = _number(row[i]) if i < len(row) else None
            if qty and qty > 0:
                product.inbound.append(Arrival(due, qty))
                count += 1
    workbook.close()
    return count


def _systeme_transit(archive: zipfile.ZipFile, info: zipfile.ZipInfo, products: dict[str, Product], as_of: date) -> int:
    workbook, sheet = _sheet(archive, info)
    rows = sheet.iter_rows(values_only=True)
    next(rows)  # merged title row
    header = next(rows)
    code_col = _find_columns(header, "Код 1с")
    category_col = _find_columns(header, "Категория 2026")
    growth_col = _find_columns(header, "Кэф. Роста")
    stock_col = _find_columns(header, "Свободный остаток")
    inbound_cols: list[tuple[int, date]] = []
    for i, value in enumerate(header):
        match = re.search(r"в пути (\d{2})\.(\d{2})", str(value or ""))
        if match:
            inbound_cols.append((i, date(as_of.year, int(match[2]), int(match[1]))))
    count = 0
    for row in rows:
        code = _code(row[code_col]) if code_col < len(row) else None
        if not code:
            continue
        product = products.setdefault(code, Product("", code, "", {}))
        category = row[category_col] if category_col < len(row) else None
        product.category = str(category).strip() if category is not None else None
        growth = _number(row[growth_col]) if growth_col < len(row) else None
        product.source_growth = growth
        stock = _number(row[stock_col]) if stock_col < len(row) else None
        if stock is not None:
            product.free_stock = stock
            product.stock_as_of = as_of
        for i, due in inbound_cols:
            qty = _number(row[i]) if i < len(row) else None
            if qty and qty > 0:
                product.inbound.append(Arrival(due, qty))
        count += 1
    workbook.close()
    return count


def load_archives(paths: list[Path], as_of: date = date(2026, 9, 22)) -> tuple[list[Product], dict]:
    """Load either or both partner ZIP files; never extract source files into Git."""
    products: dict[tuple[str, str], Product] = {}
    report = {"as_of": as_of.isoformat(), "archives": [], "files": [], "warnings": []}
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
        supplier = "IEK" if "iek" in path.stem.lower() or "иэк" in path.stem.lower() else "Systeme electric"
        local: dict[str, Product] = {}
        with zipfile.ZipFile(path) as archive:
            members = {(_decoded_name(info)).split("/")[-1]: info for info in archive.infolist() if info.filename.endswith(".xlsx")}
            def find(prefix: str) -> zipfile.ZipInfo:
                matches = [info for name, info in members.items() if name.lower().startswith(prefix.lower())]
                if len(matches) != 1:
                    raise ValueError(f"Expected one {prefix} workbook in {path.name}; got {len(matches)}")
                return matches[0]
            tasks = [
                ("monthly_sales", find("Ежемесячные продажи"), lambda info: _monthly(archive, info, local, "sales")),
                ("monthly_stock", find("Ежемесячные остатки"), lambda info: _monthly(archive, info, local, "stock")),
                ("moq", find("MOQ"), lambda info: _moq(archive, info, local)),
                ("detail_sales", find("Динамика продаж"), lambda info: _detail(archive, info, local)),
            ]
            if supplier == "IEK":
                tasks.append(("inbound", find("Путь ИЭК"), lambda info: _iek_transit(archive, info, local)))
            else:
                tasks.append(("inbound", find("Товар в пути"), lambda info: _systeme_transit(archive, info, local, as_of)))
            for kind, info, read in tasks:
                count = read(info)
                report["files"].append({"archive": path.name, "kind": kind, "rows": count})
            seasonal_workbook, seasonal_sheet = _sheet(archive, find("Сезонность"))
            report["files"].append({"archive": path.name, "kind": "seasonality_reference",
                                    "rows": seasonal_sheet.max_row, "used_in_calculation": False})
            seasonal_workbook.close()
        for code, product in local.items():
            product.supplier = supplier
            if not product.name:
                product.name = code
            if not product.monthly_sales:
                continue
            products[(supplier, code)] = product
        report["archives"].append(path.name)
        if supplier == "IEK":
            report["warnings"].append("IEK: нет подтверждённого текущего свободного остатка; количество заказа не рассчитывается")
    report["products"] = len(products)
    report["ready_stock"] = sum(p.stock_as_of == as_of for p in products.values())
    report["warnings"].extend([
        "Нет обезличенного ID клиента: крупные заказы определяются по документу и месячному всплеску",
        "Нет точных периодов stockout: нулевые начальные остатки используются только как оценка",
        "Нет справочника сроков поставки: горизонт расчёта задаётся пользователем",
    ])
    return list(products.values()), report
